import os
import time
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field

load_dotenv()
my_api_key=os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("API key kaha hai bhai")

client=Groq(api_key=my_api_key)
model = "openai/gpt-oss-120b"


job_description="""
Technology Development Program
Software Engineering Associate – 2027
Join our team of technology professionals who are deeply engaged in determining how we can change and
improve health care. We encourage you to apply today. Find out how Optum’s Technology Development
Program can lead to Caring. Connecting. Growing Together.
About Optum
Optum, part of the UnitedHealth Group family of businesses, is a leading information and technologyenabled health services company serving more than 132 million consumers, delivering the right care when
and how it’s needed. We are dedicated to evolving health care so everyone can have the opportunity to live
their healthiest life.
What is it like to work in Technology at Optum:
Combine two of the fastest-growing fields on the planet - technology and health care - with a culture of
performance, collaboration, and opportunity and this is what you get: leading-edge solutions that improve
the lives of millions. At Optum, our approach to technology isn’t about designing another app to fix health
care. It’s about connecting patients to their physician and connecting a disconnected health system so that it
works better for the people it was created to serve.
Our global technology team consists of more than 35000 innovators from around the world passionate about
using technology to improve the lives of people and communities. We are modernizing the health care
system through modern and innovative technologies, AI, comprehensive clinical expertise, and extensive
experience in improving operational and clinical performance.
As part of our team, you will be provided opportunities to channel your curiosity into initiative, working
alongside distinguished engineers and developers pursuing innovations that improve the health care 
experience for everyone. We encourage you to dive-in, explore and ask questions about creative ways to
use technology to help make the health care system work like an actual system.
Technology Development Program:
The Technology Development Program (TDP) at Optum begins with a 4-week bootcamp, followed by a 12-
month development experience that includes two 6-month rotations and concludes with a grad placement.
As a TDP Associate, you’ll gain invaluable experience across a variety of domains, benefit from mentorship
by technology leaders, and engage in networking and collaboration opportunities across the technology
organization. The program’s cohort-based structure fosters a more personalized employment experience,
helping to build community and drive engagement beyond that of a traditional hire.
Key Information Regarding Software Engineering:
• The program begins with a four-week in-office onboarding bootcamp focused on building
foundational skills, fostering community, and preparing you for success.
• After bootcamp, TDP Associates join their assigned teams and begin contributing to meaningful, realworld projects aligned with business needs and personal interests
• For 2027, TDP will be offering only Software Engineering role. The role includes, but is not limited
to analysis, design, coding, engineering, testing, debugging, standards, methods, tools analysis,
documentation, research and development, maintenance, new development, operations, and
delivery. With every role in the company, each position has a requirement for building quality into
every output. This also includes evaluating new tools, new techniques, strategies; Automation of
common tasks; build of common utilities to drive organizational efficiency with a passion around
technology and solutions and influence of thought and leadership on future capabilities and
opportunities to apply technology in new and innovative way
• TDP Associates are expected to follow the hybrid work policy of their assigned office location after
bootcamp, there is no full-time telecommuting role
• The start date for the 2027 cohort is July 27, 2027.
• Office locations* include Gurugram, Hyderabad, Bengaluru and Chennai.
• Sponsorship is not available for this position.
You’ll be rewarded and recognized for your performance in an environment that will challenge you and give
you clear direction on what it takes to succeed in your role as well as provide development for potential
future roles.
* You will be assigned location basis business requirement.
Required Qualifications:
• Degree in BE/ BTECH/ MCA in Computer Science, Computer Engineering, Software Engineering,
Data Science, IT & other technology related fields; CGPA- 7.0 or above, No active backlogs
• Currently in final year of obtaining a bachelor's or master's degree from an accredited
college/university in majors stated above
• Basic understanding of Artificial Intelligence (AI), Machine Learning (ML), Large Language Models
(LLMs), with ability to leverage AI tools to enhance productivity and problem-solving
• Proficient in one or more coding languages (Example: Python, Java, C, SQL, .NET)
• Must be eligible to work in India without company sponsorship, now or in the future, for employmentbased work authorization.
Preferred Qualifications:
• Involvement in any additional technical activities (Robotics, Comp. Science/Tech clubs, Tech or
Engineering side projects, GitHub, etc.) and/or tech related employment experience
• Cloud Technologies (AWS, Azure, and/or GCP)
• Demonstrated capability to accomplish the following:
▪ Apply AI tools to solve business problems
▪ Use APIs and frameworks to create scalable AI solutions
▪ Build and optimize prompts for analysis, reporting, and automation
▪ Develop and deploy machine learning and AI models
▪ Integrate AI into workflows (e.g., reporting, dashboards, operations processes)
▪ Analyze data using AI-enabled tools and generate insights
▪ Collaborate with teams to identify opportunities for AI-driven improvements
▪ Ensure responsible and compliant use of AI solutions
• Proven ability to collaborate effectively and communicate clearly within technical teams and crossfunctional environments
• Experience with software engineering tools such as source code control systems (i.e. GIT, GitHub),
IDE, Continuous Delivery and deployment pipeline, etc.
• Understanding of Agile Frameworks, methodologies, roles, and tools (Kanban, Scrum, SaFE, LeSS,
etc)

"""
class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibilities: list[str]

jobd_schema = JobD.model_json_schema()

system_prompt = f"""
You are an expert HR assistant.

Your job is to analyze job descriptions and extract
structured information from them.

Return ONLY valid JSON matching this schema:

{jobd_schema}
IMPORTANT:
Do NOT return the schema itself.
Do NOT return fields like "properties", "title" or "type".
Fill the schema with actual information extracted from the job description.

If minimum experience is not mentioned, return null.
If information for a list is missing, return an empty list.
Do not invent information.
"""

user_prompt = f"""
Analyze the following job description:

{job_description}
"""
message_system={
    "role" : "system",
    "content" : system_prompt
}
message_user={
    "role" : "user",
    "content" : user_prompt
}
response_format={
    "type" : "json_object"
}


messages=[message_system, message_user]

response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)


answer=response.choices[0].message.content

raw_json=answer
# print(raw_json)



import json
job_data=json.loads(raw_json)

job = JobD(**job_data)

print(job.minimum_experience)
print(job.education_requirements)



#parse real
class MatchResult(BaseModel):
    score: float
    details: dict
class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


resume_schema = Resume.model_json_schema()
def final_score(job,resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    You are an HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Return JSON matching this schema:

    {match_schema}

    Give me:

    1. Candidate name
    2. Matching skills
    3. Missing important skills
    4. Whether experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.
    """
    message={
        "role": "user",
        "content" : prompt
    }
    messages=[message]
    response_format={
        "type": "json_object"
    }
    response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)
def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt = f"""
    Parse the following resume:

    {resume_text}
    """
    message_system={
        "role" : "system",
        "content" : system_prompt
    }
    message_user={
        "role" : "user",
        "content" : user_prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume


from pypdf import PdfReader
from docx import Document
def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None



# lets do it now
resume_folder = Path("resumes")
all_results=[]
for file_path in resume_folder.iterdir():
    #C:\Users\Pratyush\padho_with_pratyush\week1\day5\resumes\abhay resume new - Abhay Singh.pdf
    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue
    print("\nProcessing:", file_path.name)
    resume_text = read_resume(file_path)
    parsed_resume=parse_resume(resume_text) # llm call1
    time.sleep(5)
    result = final_score(job, parsed_resume) #llm caLL2
    #score and details
    #acount chtgpt
    # request bhejna shhur krega millions
    #chattgot server jam ho jayega
    time.sleep(5)
    print("Score:", result.score)
    all_results.append({
        "name": parsed_resume.name,
        "score": result.score,
        "details": result.details
    })
all_results.sort(
    key=lambda candidate: candidate["score"],
    reverse=True
)
top_2 = all_results[:2]
worst_2 = all_results[-2:]


print("TOP 2 CANDIDATES")
for candidate in top_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )

    print(candidate["details"])

print("LOWEST 2 CANDIDATES")
for candidate in worst_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )
    print(candidate["details"])